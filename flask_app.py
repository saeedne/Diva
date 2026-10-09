"""Flask application for searching public Divar listings."""

from __future__ import annotations

import csv
import math
import re
import threading
import time
from copy import deepcopy
from io import BytesIO, StringIO
from pathlib import Path

from flask import Flask, Response, jsonify, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from divi import (DivarError, _parse_divar_url, _price_number, get_ad_details,
                       get_category_list, get_city_list, get_district_list, search_divar)

BASE_DIR = Path(__file__).resolve().parent
app = Flask(__name__)
_SEARCH_CACHE = {}
_SEARCH_CACHE_LOCK = threading.Lock()


BASE_EXPORT_FIELDS = [
    ("شهر", "fixed:city"), ("محله", "fixed:district"),
    ("قیمت هر متر", "fixed:price_per_meter"), ("قیمت", "fixed:price"),
    ("متراژ", "fixed:area"), ("متراژ زمین", "fixed:land_area"), ("عنوان", "fixed:title"),
]


def normalize_label(value):
    return re.sub(r"[\s‌\u200c؟?]+", "", str(value or "")).replace("ي", "ی").replace("ك", "ک").casefold()


def find_attribute(ad, *labels):
    attributes = ad.get("attributes") or {}
    normalized = {normalize_label(key): value for key, value in attributes.items()}
    for label in labels:
        value = normalized.get(normalize_label(label))
        if value not in (None, ""):
            return value
    return None


def attribute_number(value):
    if value in (None, ""):
        return None
    text = str(value).translate(str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789"))
    text = text.replace("٬", ",").replace("٫", ".")
    match = re.search(r"\d+(?:[,.]\d+)?", text)
    if not match:
        return value
    try:
        number = float(match.group().replace(",", "."))
        return int(number) if number.is_integer() else number
    except ValueError:
        return value


def dynamic_columns(ads):
    columns = list(BASE_EXPORT_FIELDS)
    reserved = {normalize_label(label) for label, _ in BASE_EXPORT_FIELDS}
    reserved.update(normalize_label(label) for label in (
        "شناسه", "دسته", "دسته‌بندی", "توضیحات کامل", "توضیحات", "قیمت کل", "قیمت کل ملک",
        "قیمت هر متر", "قیمت هر متر مربع", "قیمت متری", "متراژ", "زیربنا", "مساحت",
        "متراژ بنا", "مساحت بنا", "متراژ زمین", "مساحت زمین", "متراژ کل زمین", "مساحت کل زمین",
        "شهر", "محله", "قیمت",
    ))
    seen = set(reserved)
    for ad in ads:
        for label in (ad.get("attributes") or {}):
            key = normalize_label(label)
            if "تصویر" in key and "همینملک" in key:
                continue
            if any(term in key for term in ("قیمتکلملک", "توضیحاتکامل", "شناسهآگهی", "دستهبندی")):
                continue
            if label and key not in seen:
                columns.append((label, f"spec:{label}"))
                seen.add(key)
        for label in (ad.get("amenities") or {}):
            key = normalize_label(label)
            if label and key not in seen:
                columns.append((label, f"amenity:{label}"))
                seen.add(key)
    columns.extend([("لینک آگهی", "fixed:url"), ("تصویر", "fixed:image_url")])
    return columns


def export_value(ad, field, csv_mode=False):
    if field.startswith("fixed:"):
        name = field[6:]
        if name == "land_area":
            value = attribute_number(find_attribute(ad, "متراژ زمین", "مساحت زمین", "متراژ کل زمین", "مساحت کل زمین"))
        elif name == "area":
            value = ad.get("area") or attribute_number(find_attribute(ad, "متراژ", "زیربنا", "مساحت", "متراژ بنا", "مساحت بنا"))
        elif name == "price_per_meter":
            value = ad.get("price_per_meter") or _price_number(
                find_attribute(ad, "قیمت هر متر", "قیمت هر متر مربع", "قیمت متری"))
        elif name == "price":
            value = ad.get("total_price") or _price_number(
                find_attribute(ad, "قیمت کل", "قیمت کل ملک", "قیمت") or ad.get("price"))
        elif name == "image_url":
            has_matching_photos = ""
            for label, answer in (ad.get("attributes") or {}).items():
                normalized = normalize_label(label)
                if "تصویر" in normalized and "همینملک" in normalized:
                    has_matching_photos = normalize_label(answer)
                    break
            affirmative = has_matching_photos.startswith(("بله", "بلی", "آری", "yes"))
            value = ad.get("image_url") if affirmative else ""
        else:
            value = ad.get(name)
        if csv_mode and name in {"price", "price_per_meter"} and isinstance(value, (int, float)):
            return f"{int(value):,}"
        return value if value is not None else ""
    if field.startswith("spec:"):
        return (ad.get("attributes") or {}).get(field[5:], "") or ""
    if field.startswith("amenity:"):
        value = (ad.get("amenities") or {}).get(field[8:])
        return "دارد" if value is True else "ندارد" if value is False else ""
    return ad.get(field, "") or ""


def fetch_ads(city_slug, category_slugs, limit, query_text="", district_ids=None,
              batch_size=None, pause_seconds=1.0):
    category_slugs = tuple(dict.fromkeys(category_slugs))
    query_text = query_text.strip()
    district_ids = tuple(sorted(set(district_ids or [])))
    key = (city_slug, category_slugs, limit, query_text, district_ids, batch_size, pause_seconds)
    now = time.monotonic()
    with _SEARCH_CACHE_LOCK:
        cached = _SEARCH_CACHE.get(key)
        if cached and now - cached[0] < 300:
            return deepcopy(cached[1])
    per_category = max(1, (limit + len(category_slugs) - 1) // len(category_slugs))
    ads = []
    seen_tokens = set()
    for category_slug in category_slugs:
        for ad in search_divar(city_slug, category_slug, per_category, query_text=query_text,
                               district_ids=list(district_ids), batch_size=batch_size,
                               pause_seconds=pause_seconds):
            if ad["token"] not in seen_tokens:
                seen_tokens.add(ad["token"])
                ads.append(ad)
                if len(ads) >= limit:
                    break
        if len(ads) >= limit:
            break
    with _SEARCH_CACHE_LOCK:
        _SEARCH_CACHE[key] = (time.monotonic(), deepcopy(ads))
    return ads


@app.errorhandler(DivarError)
def handle_divar_error(error: DivarError):
    return jsonify({"detail": error.detail}), error.status_code


@app.get("/")
def home():
    return send_file(BASE_DIR / "index.html")


@app.get("/health")
def health():
    return jsonify({"ok": True, "service": "unofficial-divar-search"})


@app.get("/catalog")
def catalog():
    return jsonify({"cities": get_city_list(), "categories": get_category_list()})


@app.get("/catalog/cities")
def city_catalog():
    return jsonify({"cities": get_city_list()})


@app.get("/catalog/districts")
def district_catalog():
    city_slug = request.args.get("city", "").strip()
    if not re.fullmatch(r"[A-Za-z0-9_-]{1,80}", city_slug):
        raise DivarError(400, "شهر معتبر نیست.")
    return jsonify({"districts": get_district_list(city_slug)})


def search_parameters():
    url = request.args.get("url", "").strip()
    raw_limit = request.args.get("limit", "50")
    try:
        limit = int(raw_limit)
    except ValueError as exc:
        raise DivarError(400, "تعداد آگهی باید عدد باشد.") from exc
    if not 1 <= limit <= 100:
        raise DivarError(400, "تعداد آگهی باید بین ۱ تا ۱۰۰ باشد.")
    city_slug, url_category = _parse_divar_url(url)
    category_slugs = request.args.getlist("category") or [url_category]
    category_slugs = list(dict.fromkeys(category_slugs))
    if len(category_slugs) > 5:
        raise DivarError(400, "حداکثر ۵ دسته‌بندی را هم‌زمان انتخاب کن.")
    if any(not re.fullmatch(r"[A-Za-z0-9_-]{1,80}", slug) for slug in category_slugs):
        raise DivarError(400, "شناسهٔ دسته‌بندی معتبر نیست.")
    query_text = request.args.get("query", "").strip()
    if len(query_text) > 200:
        raise DivarError(400, "عبارت جست‌وجو نباید بیشتر از ۲۰۰ نویسه باشد.")
    district_ids = list(dict.fromkeys(request.args.getlist("district")))
    if len(district_ids) > 20 or any(not re.fullmatch(r"\d{1,12}", item) for item in district_ids):
        raise DivarError(400, "شناسهٔ محله معتبر نیست یا تعداد محله‌ها بیشتر از حد مجاز است.")
    return city_slug, category_slugs, limit, query_text, district_ids


@app.get("/ads")
def get_ads():
    city_slug, category_slugs, limit, query_text, district_ids = search_parameters()
    ads = fetch_ads(city_slug, category_slugs, limit, query_text, district_ids)
    columns = dynamic_columns(ads)
    return jsonify({"source": "divar.ir", "city": city_slug, "categories": category_slugs,
                    "category": category_slugs[0],
                    "count": len(ads), "columns": [{"label": label, "field": field}
                                                      for label, field in columns],
                    "rows": [[export_value(ad, field) for _, field in columns] for ad in ads],
                    "ads": ads})


@app.get("/ads.csv")
def download_ads_csv():
    city_slug, category_slugs, limit, query_text, district_ids = search_parameters()
    ads = fetch_ads(city_slug, category_slugs, limit, query_text, district_ids)
    output = StringIO(newline="")
    writer = csv.writer(output)
    columns = dynamic_columns(ads)
    writer.writerow([label for label, _ in columns])
    writer.writerows([[export_value(ad, field, csv_mode=True) for _, field in columns] for ad in ads])
    return Response("\ufeff" + output.getvalue(), mimetype="text/csv; charset=utf-8",
                    headers={"Content-Disposition": 'attachment; filename="divar_ads.csv"'})


@app.get("/ads.xlsx")
def download_ads_xlsx():
    city_slug, category_slugs, limit, query_text, district_ids = search_parameters()
    ads = fetch_ads(city_slug, category_slugs, limit, query_text, district_ids)
    columns = dynamic_columns(ads)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "آگهی‌های دیوار"
    sheet.sheet_view.rightToLeft = True
    sheet.append([label for label, _ in columns])
    for ad in ads:
        sheet.append([export_value(ad, field) for _, field in columns])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="5941A9")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, (label, _) in enumerate(columns, 1):
        width = min(48, max(16, len(label) * 2))
        sheet.column_dimensions[get_column_letter(index)].width = width
    price_columns = {index for index, (_, field) in enumerate(columns, 1)
                     if field in {"fixed:price", "fixed:price_per_meter"}}
    for row_index, row in enumerate(sheet.iter_rows(min_row=2), start=2):
        row_fill = PatternFill("solid", fgColor="F4F1FA" if row_index % 2 == 0 else "FFFFFF")
        for column_index, cell in enumerate(row, start=1):
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.fill = row_fill
            if column_index in price_columns and isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0"
    hyperlink_columns = {index: field for index, (_, field) in enumerate(columns, 1)
                         if field in {"fixed:url", "fixed:image_url"}}
    for column_index, field in hyperlink_columns.items():
        for row_index in range(2, sheet.max_row + 1):
            cell = sheet.cell(row_index, column_index)
            if cell.value:
                cell.hyperlink = str(cell.value)
                cell.value = "مشاهده آگهی" if field == "fixed:url" else "مشاهده تصویر"
                cell.font = Font(color="0563C1", underline="single")
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return send_file(stream, as_attachment=True, download_name=f"divar_{city_slug}_{category_slugs[0]}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def _numeric_feature(value):
    if value in (None, ""):
        return None
    words = {"بدوناتاق": 0, "یک": 1, "یه": 1, "دو": 2, "سه": 3, "چهار": 4, "بیشتر": 5}
    normalized = normalize_label(value)
    if normalized in words:
        return words[normalized]
    number = attribute_number(value)
    return number if isinstance(number, (int, float)) else None


def _ad_property_values(ad):
    area = ad.get("area") or attribute_number(find_attribute(ad, "متراژ", "زیربنا", "مساحت", "متراژ بنا"))
    land_area = attribute_number(find_attribute(ad, "متراژ زمین", "مساحت زمین", "متراژ کل زمین"))
    price = ad.get("total_price") or _price_number(ad.get("price"))
    if not price:
        price = _price_number(find_attribute(ad, "قیمت", "قیمت کل", "قیمت کل ملک"))
    year = _numeric_feature(ad.get("year_built") or find_attribute(ad, "سال ساخت", "سال بنا", "ساخت"))
    rooms = _numeric_feature(ad.get("rooms") or find_attribute(ad, "تعداد خواب", "اتاق", "خواب"))
    floor = _numeric_feature(ad.get("floor") or find_attribute(ad, "طبقه"))
    return area, land_area, price, year, rooms, floor


def _weighted_quantile(values, fraction):
    ordered = sorted(values, key=lambda pair: pair[0])
    total_weight = sum(weight for _, weight in ordered)
    cutoff = total_weight * fraction
    running = 0.0
    for value, weight in ordered:
        running += weight
        if running >= cutoff:
            return value
    return ordered[-1][0]


def _amenity_match(ad, amenity_key):
    aliases = {
        "elevator": ("آسانسور", "elevator"),
        "parking": ("پارکینگ", "parking"),
        "storage": ("انباری", "storage"),
        "balcony": ("بالکن", "تراس", "balcony"),
    }
    needles = tuple(normalize_label(item) for item in aliases.get(amenity_key, (amenity_key,)))
    for label, value in (ad.get("amenities") or {}).items():
        normalized = normalize_label(label)
        if any(needle in normalized for needle in needles):
            return bool(value)
    return None


def _property_similarity(subject, ad):
    area, land_area, _, year, rooms, floor = _ad_property_values(ad)
    if not area or area <= 0:
        return 0.0
    dimensions = []
    area_distance = abs(math.log(max(area, 1) / subject["area"]))
    dimensions.append((0.55, math.exp(-0.5 * (area_distance / 0.35) ** 2)))
    for weight, target, actual, scale in (
        (0.18, subject.get("rooms"), rooms, 2.0),
        (0.12, subject.get("year_built"), year, 15.0),
        (0.05, subject.get("floor"), floor, 4.0),
    ):
        if target is not None and actual is not None:
            dimensions.append((weight, math.exp(-0.5 * (abs(float(actual) - float(target)) / scale) ** 2)))
    if subject.get("land_area") and land_area:
        distance = abs(math.log(max(land_area, 1) / subject["land_area"]))
        dimensions.append((0.05, math.exp(-0.5 * (distance / 0.5) ** 2)))
    feature_scores = []
    for key, target in subject["amenities"].items():
        actual = _amenity_match(ad, key)
        if actual is not None:
            feature_scores.append(float(actual == target))
    if feature_scores:
        dimensions.append((0.05, sum(feature_scores) / len(feature_scores)))
    return sum(weight * score for weight, score in dimensions) / sum(weight for weight, _ in dimensions)


@app.get("/estimate")
def estimate_page():
    return send_file(BASE_DIR / "estimate.html")


@app.post("/estimate/property")
def estimate_property():
    data = request.get_json(silent=True) or {}
    if not isinstance(data, dict):
        raise DivarError(400, "اطلاعات ارسالی معتبر نیست.")
    city_slug = str(data.get("city", "")).strip()
    if not re.fullmatch(r"[A-Za-z0-9_-]{1,80}", city_slug):
        raise DivarError(400, "شهر را انتخاب کن.")
    try:
        area = float(data.get("area", 0))
    except (TypeError, ValueError) as exc:
        raise DivarError(400, "متراژ بنا باید عدد باشد.") from exc
    if not 10 <= area <= 10000:
        raise DivarError(400, "متراژ بنا باید بین ۱۰ تا ۱۰٬۰۰۰ متر باشد.")
    raw_ad_count = data.get("ad_count", 500)
    try:
        ad_count = int(raw_ad_count)
    except (TypeError, ValueError) as exc:
        raise DivarError(400, "تعداد آگهی باید عدد صحیح باشد.") from exc
    if isinstance(raw_ad_count, bool) or str(raw_ad_count).strip() != str(ad_count) or not 1 <= ad_count <= 500:
        raise DivarError(400, "تعداد آگهی برای بررسی باید بین ۱ تا ۵۰۰ باشد.")
    raw_districts = data.get("districts") or []
    if not isinstance(raw_districts, list):
        raise DivarError(400, "فهرست محله‌ها معتبر نیست.")
    district_ids = list(dict.fromkeys(str(item).strip() for item in raw_districts))
    if len(district_ids) > 20 or any(not re.fullmatch(r"\d{1,12}", item) for item in district_ids):
        raise DivarError(400, "محله‌ها معتبر نیستند یا بیش از ۲۰ محله انتخاب شده است.")
    if district_ids:
        valid_ids = {item["id"] for item in get_district_list(city_slug)}
        if not set(district_ids).issubset(valid_ids):
            raise DivarError(400, "یکی از محله‌های انتخاب‌شده در فهرست این شهر نیست.")

    transaction = data.get("transaction")
    if transaction not in {"sale", "rent"}:
        raise DivarError(400, "نوع معامله را فروش یا اجاره انتخاب کن.")
    property_types = data.get("property_types") or []
    if not isinstance(property_types, list) or not property_types:
        raise DivarError(400, "حداقل یک دستهٔ ملکی انتخاب کن.")
    category_map = {
        "sale": {"apartment": "apartment-sell", "residential": "buy-residential",
                 "villa": "house-villa-sell"},
        "rent": {"apartment": "apartment-rent", "residential": "rent-residential",
                 "villa": "house-villa-rent"},
    }
    if any(item not in category_map[transaction] for item in property_types):
        raise DivarError(400, "یکی از دسته‌های ملکی انتخاب‌شده معتبر نیست.")
    category_slugs = list(dict.fromkeys(category_map[transaction][item] for item in property_types))

    subject = {"area": area, "amenities": {}}
    ranges = {"land_area": (0, 100000), "year_built": (1300, 1500), "floor": (-5, 100), "rooms": (0, 20)}
    for key in ("land_area", "year_built", "floor", "rooms"):
        raw = data.get(key)
        if raw not in (None, ""):
            try:
                value = float(raw)
            except (TypeError, ValueError) as exc:
                raise DivarError(400, f"مقدار {key} باید عدد باشد.") from exc
            if not ranges[key][0] <= value <= ranges[key][1]:
                raise DivarError(400, f"مقدار {key} خارج از محدودهٔ مجاز است.")
            subject[key] = value
    for key in ("elevator", "parking", "storage", "balcony"):
        if isinstance(data.get(key), bool):
            subject["amenities"][key] = data[key]

    ads = fetch_ads(city_slug, category_slugs, ad_count, district_ids=district_ids,
                    batch_size=50, pause_seconds=1.0)
    comparable_values = []
    for ad in ads:
        ad_area, ad_land_area, ad_price, *_ = _ad_property_values(ad)
        if not ad_area or not ad_price or ad_area <= 0 or ad_price <= 0:
            continue
        area_ratio = ad_area / area
        if not 0.4 <= area_ratio <= 2.5:
            continue
        score = _property_similarity(subject, ad)
        if score <= 0:
            continue
        normalized_price = ad_price * area / ad_area
        comparable_values.append({
            "ad": ad, "score": score, "area": ad_area, "land_area": ad_land_area,
            "asking_price": ad_price, "adjusted_price": normalized_price,
        })
    comparable_values.sort(key=lambda item: item["score"], reverse=True)
    comparables = comparable_values[:25]
    if not comparables:
        raise DivarError(422, "آگهی مشابه با اطلاعات کافی برای تخمین در این شهر یا محله پیدا نشد.")

    weighted_prices = [(item["adjusted_price"], max(0.01, item["score"] ** 2)) for item in comparables]
    estimate = int(_weighted_quantile(weighted_prices, 0.5))
    low = int(_weighted_quantile(weighted_prices, 0.2))
    high = int(_weighted_quantile(weighted_prices, 0.8))
    if low == high:
        low, high = int(estimate * 0.9), int(estimate * 1.1)
    count = len(comparables)
    spread = (high - low) / estimate if estimate else 1
    confidence = "زیاد" if count >= 15 and spread < 0.4 else "متوسط" if count >= 5 and spread < 0.7 else "کم"
    result_comparables = []
    for item in comparables[:12]:
        ad = item["ad"]
        result_comparables.append({
            "title": ad.get("title") or "آگهی ملک", "district": ad.get("district") or "",
            "area": item["area"], "price": item["asking_price"], "url": ad.get("url"),
            "price_per_meter": int(item["asking_price"] / item["area"]),
            "similarity": round(item["score"] * 100),
        })
    method_features = "متراژ بنا و میزان شباهت در سال ساخت، تعداد اتاق، طبقه، متراژ زمین و امکانات"
    return jsonify({
        "estimate": estimate, "estimate_per_meter": int(estimate / area),
        "low": low, "high": high, "confidence": confidence,
        "comparable_count": count, "fetched_count": len(ads), "comparables": result_comparables,
        "transaction": transaction,
        "method": f"میانهٔ وزنی موارد مشابه با درنظرگرفتن {method_features}",
    })


@app.get("/ads/<token>")
def get_ad(token: str):
    return jsonify(get_ad_details(token))


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=8000, debug=False)
